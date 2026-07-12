#!/usr/bin/env python3
"""Extract native text, headings, tables, and images from pptx slides.
See SKILL.md."""

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "lib"))
import elements as elements_lib  # noqa: E402
import paths  # noqa: E402

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def shape_paragraphs(shape) -> list[str]:
    if not shape.has_text_frame:
        return []
    paragraphs = []
    for para in shape.text_frame.paragraphs:
        text = "".join(run.text for run in para.runs).strip()
        if text:
            paragraphs.append(text)
    return paragraphs


def table_rows(shape) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in shape.table.rows]


def extract_slide(slide, page_number: int, assets_dir: Path) -> tuple[list[dict], list[dict]]:
    """Returns (body_elements, image_elements) for one slide."""
    title_shape = slide.shapes.title
    title_shape_id = title_shape.shape_id if title_shape is not None else None
    body_elements = []
    image_elements = []
    bitmap_idx = 0

    for shape in slide.shapes:
        y = shape.top if shape.top is not None else 0

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            bitmap_idx += 1
            image = shape.image
            out_name = f"page{page_number}_bitmap{bitmap_idx}.{image.ext}"
            (assets_dir / out_name).write_bytes(image.blob)
            image_elements.append({"type": "image", "kind": "bitmap", "asset": f"assets/{out_name}", "caption": ""})
            continue

        if shape.has_table:
            body_elements.append({"type": "table", "rows": table_rows(shape), "y": y})
            continue

        # `slide.shapes.title` and the shapes yielded by iterating
        # `slide.shapes` are separate wrapper objects for the same XML
        # element (python-pptx builds a fresh proxy per access), so identity
        # comparison never matches — compare the stable shape_id instead.
        if title_shape_id is not None and shape.shape_id == title_shape_id:
            title_text = " ".join(shape_paragraphs(shape)).strip()
            if title_text:
                body_elements.append({"type": "heading", "level": 1, "text": title_text, "y": y})
            continue

        for para_text in shape_paragraphs(shape):
            body_elements.append({"type": "paragraph", "text": para_text, "y": y})

    body_elements.sort(key=lambda e: e["y"])
    for el in body_elements:
        del el["y"]

    if slide.has_notes_slide:
        notes_text = slide.notes_slide.notes_text_frame.text.strip()
        if notes_text:
            body_elements.append({"type": "paragraph", "text": f"Speaker notes: {notes_text}"})

    return body_elements, image_elements


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--doc", required=True)
    parser.add_argument("--pages", required=True, help="comma-separated 1-indexed slide numbers")
    args = parser.parse_args()

    page_numbers = [int(p) for p in args.pages.split(",") if p.strip()]
    pptx_path = paths.input_file(args.doc)
    if pptx_path is None or pptx_path.suffix != ".pptx":
        print(f"error: input/{args.doc}.pptx not found", file=sys.stderr)
        sys.exit(1)

    presentation = Presentation(pptx_path)
    assets_dir = paths.assets_dir(args.doc)
    assets_dir.mkdir(parents=True, exist_ok=True)

    for page_number in page_numbers:
        slide = presentation.slides[page_number - 1]
        body_elements, image_elements = extract_slide(slide, page_number, assets_dir)

        elements_lib.write_shard(paths.shard_path(args.doc, page_number, "text"), page_number, body_elements)
        elements_lib.write_shard(paths.shard_path(args.doc, page_number, "image"), page_number, image_elements)

    print(f"extracted pptx slides {page_numbers}")


if __name__ == "__main__":
    main()
