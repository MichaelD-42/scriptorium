#!/usr/bin/env python3
"""Classify each sheet of a workbook into an extraction tier (always
"text", xlsx is digital-native) and the document into a loop size
(tight|loose). See SKILL.md for the schema."""

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "lib"))
import paths  # noqa: E402

import openpyxl


def sheet_has_data(ws) -> bool:
    for row in ws.iter_rows(values_only=True):
        if any(cell is not None for cell in row):
            return True
    return False


def classify_sheet(ws, page_number: int) -> dict:
    return {
        "page_number": page_number,
        "tier": "text",
        "has_table": sheet_has_data(ws),
        "image_count": len(ws._images),
        "has_chart": bool(ws._charts),
        "reason": "native worksheet content (digital-native, no OCR needed)",
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--doc", required=True, help="document name (without .xlsx)")
    args = parser.parse_args()

    xlsx_path = paths.input_file(args.doc)
    if xlsx_path is None or xlsx_path.suffix != ".xlsx":
        print(f"error: input/{args.doc}.xlsx not found", file=sys.stderr)
        sys.exit(1)

    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    pages = [classify_sheet(workbook[name], i) for i, name in enumerate(workbook.sheetnames, start=1)]

    loop_size = "tight" if any(p["image_count"] > 0 or p["has_chart"] for p in pages) else "loose"
    result = {
        "doc": args.doc,
        "page_count": len(pages),
        "pages": pages,
        "loop_size": loop_size,
    }

    out_path = paths.triage_json(args.doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(result, indent=2), encoding="utf-8", newline="")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
