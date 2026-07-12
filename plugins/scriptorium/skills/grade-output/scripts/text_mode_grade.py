#!/usr/bin/env python3
"""Deterministic structural grading for formats with no rendered page
(docx, xlsx). See text-rubric.md — this replaces the `grader` subagent's
visual judgment with a script that re-reads the source file directly and
compares it against elements.json, since there's no pixel ground truth for
a subagent to look at.

Writes the same per-page shard shape write_grade_shard.py does
({"page_number", "score", "issues"}), so merge_grades.py needs no changes
regardless of whether shards came from a script or a subagent.
"""

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "lib"))
import docx_pages  # noqa: E402
import elements as elements_lib  # noqa: E402
import paths  # noqa: E402

MIN_UNIT_LEN = 15  # shorter fragments are too generic to reliably match/mismatch
GENERIC_CAPTIONS = {"image", "figure", "picture", ""}
SCORE_PENALTY_PER_ISSUE_TYPE = 0.3


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip().lower())


def output_page_text(page: dict) -> str:
    parts = []
    for el in page["elements"]:
        if el["type"] in ("heading", "paragraph"):
            parts.append(el["text"])
        elif el["type"] == "table":
            parts.extend(cell for row in el["rows"] for cell in row)
    return normalize(" ".join(parts))


def score_from_issues(issues: list[str]) -> float:
    return round(max(0.0, 1.0 - SCORE_PENALTY_PER_ISSUE_TYPE * len(issues)), 3) if issues else 1.0


def write_grade_shard(doc: str, page_number: int, score: float, issues: list[str]) -> None:
    shard_path = paths.grade_shard_path(doc, page_number)
    shard_path.parent.mkdir(parents=True, exist_ok=True)
    shard_path.write_text(json.dumps({"page_number": page_number, "score": score, "issues": issues}, indent=2))
    print(f"page {page_number}: score={score} issues={issues}")


# --- docx -------------------------------------------------------------------

def docx_check_dropped_text(blocks: list, output_text: str) -> bool:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for block in blocks:
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if len(text) >= MIN_UNIT_LEN and normalize(text) not in output_text:
                return True
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if len(text) >= MIN_UNIT_LEN and normalize(text) not in output_text:
                        return True
    return False


def docx_check_table_corruption(blocks: list, output_elements: list) -> bool:
    from docx.table import Table

    source_tables = [b for b in blocks if isinstance(b, Table)]
    output_tables = [el for el in output_elements if el["type"] == "table"]
    if len(source_tables) != len(output_tables):
        return True
    for src, out in zip(source_tables, output_tables):
        src_rows = len(src.rows)
        src_cols = len(src.columns)
        out_rows = len(out["rows"])
        out_cols = max((len(r) for r in out["rows"]), default=0)
        if src_rows != out_rows or src_cols != out_cols:
            return True
    return False


def docx_check_images(blocks: list, output_elements: list, document) -> tuple[bool, bool]:
    from docx.text.paragraph import Paragraph

    source_image_count = sum(
        len(docx_pages.paragraph_images(b, document)) for b in blocks if isinstance(b, Paragraph)
    )
    output_images = [el for el in output_elements if el["type"] == "image"]
    missing_image = source_image_count > len(output_images)
    bad_caption = any((el.get("caption") or "").strip().lower() in GENERIC_CAPTIONS for el in output_images)
    return missing_image, bad_caption


def docx_check_wrong_heading_level(blocks: list, output_elements: list) -> bool:
    from docx.text.paragraph import Paragraph

    output_headings = {normalize(el["text"]): el["level"] for el in output_elements if el["type"] == "heading"}
    for block in blocks:
        if not isinstance(block, Paragraph):
            continue
        level = docx_pages.heading_level(block)
        if level is None:
            continue
        text = block.text.strip()
        if not text:
            continue
        expected_level = min(level, 3)
        actual_level = output_headings.get(normalize(text))
        if actual_level != expected_level:
            return True
    return False


def grade_docx_page(blocks: list, page: dict, document) -> tuple[float, list[str]]:
    output_elements = page["elements"]
    output_text = output_page_text(page)

    issues = []
    if docx_check_dropped_text(blocks, output_text):
        issues.append("dropped_text")
    if docx_check_table_corruption(blocks, output_elements):
        issues.append("table_corruption")
    missing_image, bad_caption = docx_check_images(blocks, output_elements, document)
    if missing_image:
        issues.append("missing_image")
    if bad_caption:
        issues.append("bad_caption")
    if docx_check_wrong_heading_level(blocks, output_elements):
        issues.append("wrong_heading_level")

    return score_from_issues(issues), issues


def grade_docx(doc: str, docx_path: Path, doc_data: dict) -> None:
    from docx import Document

    document = Document(docx_path)
    source_pages = docx_pages.split_pages(document)

    for page_number, blocks in enumerate(source_pages, start=1):
        page = doc_data["pages"].get(page_number, {"elements": []})
        score, issues = grade_docx_page(blocks, page, document)
        write_grade_shard(doc, page_number, score, issues)


# --- xlsx -------------------------------------------------------------------

def xlsx_check_dropped_text(ws, output_text: str) -> bool:
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is None:
                continue
            text = str(cell).strip()
            if len(text) >= MIN_UNIT_LEN and normalize(text) not in output_text:
                return True
    return False


def xlsx_check_table_corruption(ws, output_elements: list) -> bool:
    has_data = any(cell is not None for row in ws.iter_rows(values_only=True) for cell in row)
    output_tables = [el for el in output_elements if el["type"] == "table"]
    if not has_data:
        return bool(output_tables)  # a blank sheet shouldn't have produced a table
    if len(output_tables) != 1:
        return True
    expected_rows = ws.max_row - ws.min_row + 1
    expected_cols = ws.max_column - ws.min_column + 1
    out_rows_list = output_tables[0]["rows"]
    if len(out_rows_list) != expected_rows:
        return True
    # Check every row's width, not just the max — a single truncated row
    # amid otherwise full-width rows wouldn't show up in a max() count.
    return any(len(row) != expected_cols for row in out_rows_list)


def xlsx_check_images(ws, output_elements: list) -> tuple[bool, bool]:
    source_image_count = len(ws._images)
    output_images = [el for el in output_elements if el["type"] == "image"]
    missing_image = source_image_count > len(output_images)
    bad_caption = any((el.get("caption") or "").strip().lower() in GENERIC_CAPTIONS for el in output_images)
    return missing_image, bad_caption


def grade_xlsx_sheet(ws, page: dict) -> tuple[float, list[str]]:
    output_elements = page["elements"]
    output_text = output_page_text(page)

    issues = []
    if xlsx_check_dropped_text(ws, output_text):
        issues.append("dropped_text")
    if xlsx_check_table_corruption(ws, output_elements):
        issues.append("table_corruption")
    missing_image, bad_caption = xlsx_check_images(ws, output_elements)
    if missing_image:
        issues.append("missing_image")
    if bad_caption:
        issues.append("bad_caption")
    # No wrong_heading_level check for xlsx: a sheet has exactly one
    # heading (its own name) and no sub-heading structure to get wrong.

    return score_from_issues(issues), issues


def grade_xlsx(doc: str, xlsx_path: Path, doc_data: dict) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    for page_number, sheet_name in enumerate(workbook.sheetnames, start=1):
        ws = workbook[sheet_name]
        page = doc_data["pages"].get(page_number, {"elements": []})
        score, issues = grade_xlsx_sheet(ws, page)
        write_grade_shard(doc, page_number, score, issues)


# --- entry point --------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--doc", required=True)
    args = parser.parse_args()

    input_path = paths.input_file(args.doc)
    if input_path is None:
        print(f"error: no input file found for doc '{args.doc}' in input/", file=sys.stderr)
        sys.exit(1)
    input_format = paths.detect_input_format(args.doc)

    elements_path = paths.elements_json(args.doc)
    if not elements_path.exists():
        print(f"error: {elements_path} not found — run merge.py first", file=sys.stderr)
        sys.exit(1)
    doc_data = elements_lib.load_doc(elements_path)

    if input_format == "docx":
        grade_docx(args.doc, input_path, doc_data)
    elif input_format == "xlsx":
        grade_xlsx(args.doc, input_path, doc_data)
    else:
        print(f"error: text_mode_grade.py has no grading path for format {input_format!r}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
