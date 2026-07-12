#!/usr/bin/env python3
"""Classify each slide of a pptx into an extraction tier (always "text",
pptx is digital-native) and the document into a loop size (tight|loose).
See SKILL.md for the schema."""

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "lib"))
import paths  # noqa: E402

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def slide_has_visual(slide) -> bool:
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return True
        if shape.has_chart:
            return True
    return False


def classify_slide(slide, page_number: int) -> dict:
    image_count = sum(1 for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return {
        "page_number": page_number,
        "tier": "text",
        "image_count": image_count,
        "reason": "native slide content (pptx is digital-native, no OCR needed)",
        "has_visual": slide_has_visual(slide),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--doc", required=True, help="document name (without .pptx)")
    args = parser.parse_args()

    pptx_path = paths.input_file(args.doc)
    if pptx_path is None or pptx_path.suffix != ".pptx":
        print(f"error: input/{args.doc}.pptx not found", file=sys.stderr)
        sys.exit(1)

    presentation = Presentation(pptx_path)
    pages = [classify_slide(slide, i) for i, slide in enumerate(presentation.slides, start=1)]

    loop_size = "tight" if any(p["has_visual"] for p in pages) else "loose"
    for p in pages:
        del p["has_visual"]
    result = {
        "doc": args.doc,
        "page_count": len(pages),
        "pages": pages,
        "loop_size": loop_size,
    }

    out_path = paths.triage_json(args.doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(result, indent=2))
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
