"""Regression tests for the 0.5.0 audit findings."""

import json
import os
import stat

from bs4 import BeautifulSoup
from docx import Document
from PIL import Image

import html_pages
import paths
from conftest import load_script, run_script

text_mode = load_script("grade-output/scripts/text_mode_grade.py", "text_mode_audit")
render = load_script("render-pages/scripts/render.py", "render_audit")


def _png(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (4, 4), "white").save(path)


class TestHtmlImagesStayInTheInputFolder:
    def test_only_images_inside_the_html_folder_are_read(self, tmp_path):
        base = tmp_path / "input"
        _png(base / "img" / "ok.png")
        _png(tmp_path / "secret.png")
        (base / "notes.png").write_text("not an image")
        soup = BeautifulSoup(
            f'<img src="img/ok.png"><img src="../secret.png"><img src="{tmp_path / "secret.png"}">'
            '<img src="notes.png">',
            "html.parser",
        )

        found = html_pages.saveable_images(soup, base)

        assert [ext for _data, ext in found] == ["png"]
        assert found[0][0] == (base / "img" / "ok.png").read_bytes()


class TestRepeatedHeadingText:
    def test_docx_same_heading_text_at_two_levels_is_not_a_level_error(self):
        d = Document()
        d.add_heading("Overview", 2)
        d.add_heading("Details", 2)
        d.add_heading("Overview", 3)
        blocks = [p for p in d.paragraphs]
        out = [
            {"type": "heading", "text": t, "level": lv}
            for t, lv in (("Overview", 2), ("Details", 2), ("Overview", 3))
        ]

        assert not text_mode.docx_check_wrong_heading_level(blocks, out)
        out[2]["level"] = 2
        assert text_mode.docx_check_wrong_heading_level(blocks, out)

    def test_html_same_heading_text_at_two_levels_is_not_a_level_error(self):
        soup = BeautifulSoup(
            "<h2>Overview</h2><h2>Details</h2><h3>Overview</h3>", "html.parser"
        )
        blocks = soup.find_all(["h2", "h3"])
        out = [
            {"type": "heading", "text": t, "level": lv}
            for t, lv in (("Overview", 2), ("Details", 2), ("Overview", 3))
        ]

        assert not text_mode.html_check_wrong_heading_level(blocks, out)
        out[2]["level"] = 2
        assert text_mode.html_check_wrong_heading_level(blocks, out)


class TestQueueInit:
    def _init(self, cwd):
        result = run_script("queue-state/scripts/state.py", "init", cwd=cwd)
        assert result.returncode == 0, result.stderr
        return json.loads((cwd / "runs" / "state.json").read_text())["docs"]

    def test_a_changed_input_is_requeued_from_a_clean_slate(self, tmp_project):
        src = tmp_project / "input" / "spec.html"
        src.write_text("<p>revision A</p>")
        self._init(tmp_project)
        state_path = tmp_project / "runs" / "state.json"
        state = json.loads(state_path.read_text())
        state["docs"]["spec"].update(status="passed", attempt=2)
        state_path.write_text(json.dumps(state))
        stale = paths.shard_path("spec", 1, "vision")
        stale.parent.mkdir(parents=True)
        stale.write_text("{}")
        paths.output_dir("spec").mkdir(parents=True)

        assert (
            self._init(tmp_project)["spec"]["status"] == "passed"
        )  # unchanged input: untouched
        assert stale.exists()

        src.write_text("<p>revision B</p>")
        entry = self._init(tmp_project)["spec"]

        assert (entry["status"], entry["attempt"]) == ("pending", 0)
        assert not paths.work_dir("spec").exists()
        assert not paths.output_dir("spec").exists()

    def test_a_new_entry_drops_work_left_from_a_deleted_queue(self, tmp_project):
        (tmp_project / "input" / "spec.html").write_text("<p>x</p>")
        stale = paths.shard_path("spec", 1, "vision")
        stale.parent.mkdir(parents=True)
        stale.write_text("{}")

        self._init(tmp_project)

        assert not stale.exists()

    def test_two_inputs_with_one_name_use_the_same_file_as_the_scripts(
        self, tmp_project
    ):
        (tmp_project / "input" / "spec.html").write_text("<p>x</p>")
        (tmp_project / "input" / "spec.pdf").write_bytes(b"%PDF-1.4")

        assert (
            self._init(tmp_project)["spec"]["input_format"]
            == paths.detect_input_format("spec")
            == "pdf"
        )


class TestPptxConversion:
    def test_soffice_runs_with_its_own_profile_and_the_pdf_appears_whole(
        self, tmp_project, monkeypatch
    ):
        """Parallel graders each convert the deck; a shared LibreOffice profile
        and a half-written PDF at the final path made that unsafe."""
        bin_dir = tmp_project / "bin"
        bin_dir.mkdir()
        log = tmp_project / "soffice-args.txt"
        fake = bin_dir / "soffice"
        fake.write_text(
            "#!/bin/sh\n"
            f'echo "$@" > {log}\n'
            'while [ "$1" != "--outdir" ]; do shift; done\n'
            'printf "%%PDF-1.4" > "$2/deck.pdf"\n'
        )
        fake.chmod(fake.stat().st_mode | stat.S_IEXEC)
        monkeypatch.setenv("PATH", f"{bin_dir}{os.pathsep}{os.environ['PATH']}")
        pptx = tmp_project / "input" / "deck.pptx"
        pptx.write_bytes(b"pptx")

        out = render.convert_pptx_to_pdf(pptx, "deck")

        args = log.read_text()
        assert "-env:UserInstallation=file://" in args
        assert f"--outdir {paths.work_dir('deck')} " not in args + " "
        assert out == paths.work_dir("deck") / "deck.pdf"
        assert out.read_bytes() == b"%PDF-1.4"
        assert sorted(p.name for p in paths.work_dir("deck").iterdir()) == ["deck.pdf"]
