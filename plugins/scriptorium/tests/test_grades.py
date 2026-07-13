"""Unit tests for the grade-output/ocr-page/extract-images scripts that
aren't covered elsewhere: text_mode_grade.py's structural checks,
merge_grades.py's arithmetic, and the three agent-landing scripts
(write_grade_shard.py, write_vision_page.py, caption_image.py)."""

import io
import json

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

import docx_pages
import html_pages
import paths
from conftest import load_script

text_mode_grade = load_script("grade-output/scripts/text_mode_grade.py", "text_mode_grade_module")
merge_grades = load_script("grade-output/scripts/merge_grades.py", "merge_grades_module")
write_grade_shard = load_script("grade-output/scripts/write_grade_shard.py", "write_grade_shard_module")
write_vision_page = load_script("ocr-page/scripts/write_vision_page.py", "write_vision_page_module")
caption_image = load_script("extract-images/scripts/caption_image.py", "caption_image_module")


# --- text_mode_grade.py: docx --------------------------------------------

@pytest.fixture
def docx_page2_blocks(docx_doc):
    document = Document(paths.input_file(docx_doc))
    return docx_pages.split_pages(document)[1], document


class TestGradeDocxPage:
    def test_faithful_extraction_scores_perfect(self, docx_page2_blocks):
        blocks, document = docx_page2_blocks
        page = {
            "elements": [
                {"type": "heading", "level": 1, "text": "Quarterly Data"},
                {"type": "table", "rows": [["Quarter", "Revenue"], ["Q1", "120"], ["Q2", "150"]]},
                {"type": "paragraph", "text": "Status indicator image below."},
                {"type": "image", "asset": "assets/page2_bitmap1.png", "caption": "a blue circle icon"},
            ]
        }
        score, issues = text_mode_grade.grade_docx_page(blocks, page, document)
        assert issues == []
        assert score == 1.0

    def test_dropped_truncated_and_missing_image_all_flagged(self, docx_page2_blocks):
        blocks, document = docx_page2_blocks
        page = {
            "elements": [
                {"type": "heading", "level": 2, "text": "Quarterly Data"},  # wrong level (source is 1)
                {"type": "table", "rows": [["Quarter", "Revenue"], ["Q1", "120"]]},  # missing a row
                # paragraph and image both dropped entirely
            ]
        }
        score, issues = text_mode_grade.grade_docx_page(blocks, page, document)
        assert set(issues) == {"dropped_text", "table_corruption", "missing_image", "wrong_heading_level"}
        assert score == 0.0


# --- text_mode_grade.py: xlsx ---------------------------------------------

class TestGradeXlsxSheet:
    def test_faithful_extraction_scores_perfect(self, xlsx_doc):
        workbook = load_workbook(paths.input_file(xlsx_doc), data_only=True)
        ws = workbook["Data"]
        page = {"elements": [
            {"type": "heading", "level": 1, "text": "Data"},
            {"type": "table", "rows": [["Quarter", "Revenue"], ["Q1", "120"], ["Q2", "150"]]},
        ]}
        score, issues = text_mode_grade.grade_xlsx_sheet(ws, page)
        assert issues == []
        assert score == 1.0

    def test_truncated_table_flagged_as_corruption(self, xlsx_doc):
        workbook = load_workbook(paths.input_file(xlsx_doc), data_only=True)
        ws = workbook["Data"]
        page = {"elements": [
            {"type": "heading", "level": 1, "text": "Data"},
            {"type": "table", "rows": [["Quarter", "Revenue"]]},  # missing 2 data rows
        ]}
        score, issues = text_mode_grade.grade_xlsx_sheet(ws, page)
        assert "table_corruption" in issues


def test_xlsx_check_dropped_text_flags_a_long_missing_cell():
    wb = Workbook()
    ws = wb.active
    ws.append(["This sentence is definitely over fifteen characters long"])
    assert text_mode_grade.xlsx_check_dropped_text(ws, output_text="") is True
    assert text_mode_grade.xlsx_check_dropped_text(
        ws, output_text="this sentence is definitely over fifteen characters long"
    ) is False


def test_xlsx_check_images_flags_generic_caption_regardless_of_source_count():
    wb = Workbook()
    ws = wb.active  # no embedded images at all
    output_elements = [{"type": "image", "asset": "a.png", "caption": "image"}]
    missing_image, bad_caption = text_mode_grade.xlsx_check_images(ws, output_elements)
    assert missing_image is False  # source has 0 images, so "more images than source" can't trigger
    assert bad_caption is True


# --- text_mode_grade.py: html ----------------------------------------------

@pytest.fixture
def html_fixture_parts(html_doc):
    html_path = paths.input_file(html_doc)
    soup = html_pages.parse(html_path)
    blocks = list(html_pages.iter_block_items(soup))
    return blocks, soup, html_path.parent


class TestGradeHtmlPage:
    def test_faithful_extraction_scores_perfect(self, html_fixture_parts):
        blocks, soup, html_dir = html_fixture_parts
        page = {"elements": [
            {"type": "heading", "level": 1, "text": "Elastic Loop HTML Sample"},
            {"type": "heading", "level": 2, "text": "Introduction"},
            {"type": "paragraph", "text": "This document exercises the HTML extraction path of the scriptorium pipeline: headings, paragraphs, a table, and an inline data-URI image."},
            {"type": "paragraph", "text": "Grading here is text-mode, since there is no rendered page image for a grader subagent to judge against."},
            {"type": "heading", "level": 2, "text": "Quarterly Results"},
            {"type": "table", "rows": [["Quarter", "Revenue", "Growth"], ["Q1", "120", "-"], ["Q2", "150", "25%"], ["Q3", "180", "20%"]]},
            {"type": "heading", "level": 2, "text": "Figure: Status Indicator"},
            {"type": "paragraph", "text": "Status indicator image below."},
            {"type": "image", "asset": "assets/page1_bitmap1.png", "caption": "a blue circle status indicator icon"},
        ]}
        score, issues = text_mode_grade.grade_html_page(blocks, page, soup, html_dir)
        assert issues == []
        assert score == 1.0

    def test_dropped_text_truncated_table_and_missing_image_all_flagged(self, html_fixture_parts):
        blocks, soup, html_dir = html_fixture_parts
        page = {"elements": [
            {"type": "heading", "level": 1, "text": "Elastic Loop HTML Sample"},
            {"type": "heading", "level": 2, "text": "Introduction"},
            # both body paragraphs dropped
            {"type": "heading", "level": 3, "text": "Quarterly Results"},  # wrong level (source is 2)
            {"type": "table", "rows": [["Quarter", "Revenue", "Growth"], ["Q1", "120", "-"], ["Q2", "150", "25%"]]},  # missing a row
            {"type": "heading", "level": 2, "text": "Figure: Status Indicator"},
            # image dropped entirely
        ]}
        score, issues = text_mode_grade.grade_html_page(blocks, page, soup, html_dir)
        assert set(issues) == {"dropped_text", "table_corruption", "missing_image", "wrong_heading_level"}
        assert score == 0.0


# --- merge_grades.py --------------------------------------------------------

def _seed_gates_report(tmp_project, doc: str, passed: bool) -> None:
    gates_path = paths.gates_report_json(doc)
    gates_path.parent.mkdir(parents=True, exist_ok=True)
    gates_path.write_text(json.dumps({"doc": doc, "passed": passed, "checks": []}))


def _seed_grade_shard(tmp_project, doc: str, page: int, score: float, issues: list[str]) -> None:
    shard_path = paths.grade_shard_path(doc, page)
    shard_path.parent.mkdir(parents=True, exist_ok=True)
    shard_path.write_text(json.dumps({"page_number": page, "score": score, "issues": issues}))


class TestMergeGrades:
    def test_overall_passed_true_when_gates_pass_and_score_meets_threshold(self, tmp_project, monkeypatch):
        _seed_gates_report(tmp_project, "doc", passed=True)
        _seed_grade_shard(tmp_project, "doc", 1, 1.0, [])
        _seed_grade_shard(tmp_project, "doc", 2, 0.7, ["dropped_text"])

        monkeypatch.setattr("sys.argv", ["merge_grades.py", "--doc", "doc"])
        merge_grades.main()

        result = json.loads(paths.grade_report_json("doc").read_text())
        assert result["rubric_verdict"]["score"] == 0.85
        assert result["rubric_verdict"]["passed"] is True
        assert result["overall_passed"] is True
        assert result["rubric_verdict"]["failure_taxonomy"] == ["dropped_text"]

    def test_overall_passed_false_when_gates_fail_even_if_rubric_passes(self, tmp_project, monkeypatch):
        _seed_gates_report(tmp_project, "doc", passed=False)
        _seed_grade_shard(tmp_project, "doc", 1, 1.0, [])

        monkeypatch.setattr("sys.argv", ["merge_grades.py", "--doc", "doc"])
        merge_grades.main()

        result = json.loads(paths.grade_report_json("doc").read_text())
        assert result["rubric_verdict"]["passed"] is True
        assert result["overall_passed"] is False

    def test_missing_gates_report_exits_with_error(self, tmp_project, monkeypatch):
        monkeypatch.setattr("sys.argv", ["merge_grades.py", "--doc", "doc"])
        with pytest.raises(SystemExit) as exc_info:
            merge_grades.main()
        assert exc_info.value.code == 1

    def test_no_grade_shards_exits_with_error(self, tmp_project, monkeypatch):
        _seed_gates_report(tmp_project, "doc", passed=True)
        monkeypatch.setattr("sys.argv", ["merge_grades.py", "--doc", "doc"])
        with pytest.raises(SystemExit) as exc_info:
            merge_grades.main()
        assert exc_info.value.code == 1


# --- write_grade_shard.py / write_vision_page.py / caption_image.py -------

def test_write_grade_shard_lands_score_and_issues(tmp_project, monkeypatch):
    monkeypatch.setattr(
        "sys.argv",
        ["write_grade_shard.py", "--doc", "doc", "--page", "3", "--score", "0.7", "--issues", "dropped_text,bad_caption"],
    )
    write_grade_shard.main()
    shard = json.loads(paths.grade_shard_path("doc", 3).read_text())
    assert shard == {"page_number": 3, "score": 0.7, "issues": ["dropped_text", "bad_caption"]}


def test_write_grade_shard_empty_issues_string_yields_empty_list(tmp_project, monkeypatch):
    monkeypatch.setattr("sys.argv", ["write_grade_shard.py", "--doc", "doc", "--page", "1", "--score", "1.0"])
    write_grade_shard.main()
    shard = json.loads(paths.grade_shard_path("doc", 1).read_text())
    assert shard["issues"] == []


def test_write_vision_page_lands_elements_from_stdin(tmp_project, monkeypatch):
    monkeypatch.setattr("sys.argv", ["write_vision_page.py", "--doc", "doc", "--page", "5"])
    monkeypatch.setattr("sys.stdin", io.StringIO(json.dumps([{"type": "paragraph", "text": "scanned text"}])))
    write_vision_page.main()

    shard = json.loads(paths.shard_path("doc", 5, "vision").read_text())
    assert shard["elements"] == [{"type": "paragraph", "text": "scanned text"}]


def test_caption_image_sets_caption_on_matching_asset(tmp_project, monkeypatch):
    import elements as elements_lib

    shard_path = paths.shard_path("doc", 2, "image")
    elements_lib.write_shard(shard_path, 2, [{"type": "image", "asset": "assets/page2_bitmap1.png", "caption": ""}])

    monkeypatch.setattr(
        "sys.argv",
        ["caption_image.py", "--doc", "doc", "--page", "2", "--asset", "assets/page2_bitmap1.png", "--caption", "a red icon"],
    )
    caption_image.main()

    shard = json.loads(shard_path.read_text())
    assert shard["elements"][0]["caption"] == "a red icon"


def test_caption_image_unknown_asset_exits_with_error(tmp_project, monkeypatch):
    import elements as elements_lib

    shard_path = paths.shard_path("doc", 2, "image")
    elements_lib.write_shard(shard_path, 2, [{"type": "image", "asset": "assets/other.png", "caption": ""}])

    monkeypatch.setattr(
        "sys.argv",
        ["caption_image.py", "--doc", "doc", "--page", "2", "--asset", "assets/page2_bitmap1.png", "--caption", "x"],
    )
    with pytest.raises(SystemExit) as exc_info:
        caption_image.main()
    assert exc_info.value.code == 1
