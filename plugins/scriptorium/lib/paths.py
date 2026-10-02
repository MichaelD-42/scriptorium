"""Conventional per-document paths, resolved against the current working
directory (the orchestrator always runs skill scripts from the project root,
i.e. ${CLAUDE_PROJECT_DIR}).

    input/<doc>.{pdf,pptx,xlsx,docx,html,png,jpg,jpeg,webp,tiff}  one supported input extension
    work/<doc>/pages/page{N}.png
    work/<doc>/shards/page{N}.{text|ocr|vision}.json   one body shard per page (highest tier wins)
    work/<doc>/shards/page{N}.image.json               independent of body tier
    work/<doc>/elements.json                           merge.py's output — the merged shards
    work/<doc>/triage.json
    work/<doc>/toc.json
    work/<doc>/gates-report.json
    output/<doc>/<doc>.{md,html,reqif,reqifz}          single-file formats (reqifz is also a zip archive)
    output/<doc>/{index.md,NN-slug.md}                 okf format (multi-file bundle)
    output/<doc>/{index.md,00-front-matter.md,NN-slug/NN.MM-slug.md}  md-tree format (split multi-file bundle, configurable depth)
    output/<doc>/assets/*.png
    output/<doc>/grade-shards/page{N}.json              one per grader batch page
    output/<doc>/grade-report.json                      merge_grades.py's output
    output/<doc>.zip                                    zip_output.py's output (sibling to output/<doc>/, not inside it)
"""

from pathlib import Path

IMAGE_EXTS = ("png", "jpg", "jpeg", "webp", "tiff")

SUPPORTED_INPUT_EXTS = ("pdf", "pptx", "xlsx", "docx", "html") + IMAGE_EXTS


def input_pdf(doc: str, root: Path = Path(".")) -> Path:
    return root / "input" / f"{doc}.pdf"


def input_file(doc: str, root: Path = Path(".")) -> Path | None:
    """Resolve input/<doc>.<ext> for whichever supported extension exists.
    Returns None if no matching file is present."""
    for ext in SUPPORTED_INPUT_EXTS:
        candidate = root / "input" / f"{doc}.{ext}"
        if candidate.exists():
            return candidate
    return None


def detect_input_format(doc: str, root: Path = Path(".")) -> str | None:
    """The input file's format (e.g. "pdf", "pptx", "image"), or None if the
    document has no input file under any supported extension. Every
    extension in IMAGE_EXTS normalizes to "image" — downstream code branches
    on one format, not five extensions."""
    path = input_file(doc, root)
    if not path:
        return None
    ext = path.suffix.lstrip(".").lower()
    return "image" if ext in IMAGE_EXTS else ext


def true_page_count(doc: str, input_format: str, root: Path = Path(".")) -> int:
    """The page count merge.py/gates.py check elements.json against.

    For pdf/pptx/xlsx this is a genuinely independent oracle read straight
    from the source file (never derived from triage.json or any other
    pipeline output, or the page_count_match gate becomes tautological) —
    a workbook's sheets, like a pptx's slides, are a real source-file
    property. docx has no such oracle to read — Word pagination is
    layout-computed and invisible to python-docx, so "page" is this
    plugin's own Heading-1 split, decided once by docx-triage. For docx,
    page_count_match therefore checks "did extraction cover every section
    triage declared", not an independent source-file property — see
    lib/docx_pages.py. html has no page concept at all — the whole file is
    always one page, a fixed rule rather than something triage decides.
    image is the same fixed rule as html: one image file is one page."""
    if input_format == "pdf":
        import fitz  # PyMuPDF

        return fitz.open(input_file(doc, root)).page_count
    if input_format == "pptx":
        from pptx import Presentation

        return len(Presentation(input_file(doc, root)).slides)
    if input_format == "xlsx":
        import openpyxl

        return len(openpyxl.load_workbook(input_file(doc, root), read_only=True).sheetnames)
    if input_format == "docx":
        import json

        triage_path = triage_json(doc, root)
        if not triage_path.exists():
            raise FileNotFoundError(f"{triage_path} not found — run docx-triage first")
        return json.loads(triage_path.read_text(encoding="utf-8"))["page_count"]
    if input_format == "html":
        return 1
    if input_format == "image":
        return 1
    raise NotImplementedError(f"true_page_count: unsupported input format {input_format!r}")


def work_dir(doc: str, root: Path = Path(".")) -> Path:
    return root / "work" / doc


def pages_dir(doc: str, root: Path = Path(".")) -> Path:
    return work_dir(doc, root) / "pages"


def page_png(doc: str, page_number: int, root: Path = Path(".")) -> Path:
    return pages_dir(doc, root) / f"page{page_number}.png"


def shards_dir(doc: str, root: Path = Path(".")) -> Path:
    return work_dir(doc, root) / "shards"


def shard_path(doc: str, page_number: int, kind: str, root: Path = Path(".")) -> Path:
    """kind is one of "text", "ocr", "vision" (body shards — mutually
    exclusive per page, highest tier wins at merge time) or "image"
    (independent of body tier)."""
    return shards_dir(doc, root) / f"page{page_number}.{kind}.json"


def elements_json(doc: str, root: Path = Path(".")) -> Path:
    return work_dir(doc, root) / "elements.json"


def triage_json(doc: str, root: Path = Path(".")) -> Path:
    return work_dir(doc, root) / "triage.json"


def toc_json(doc: str, root: Path = Path(".")) -> Path:
    return work_dir(doc, root) / "toc.json"


def gates_report_json(doc: str, root: Path = Path(".")) -> Path:
    return work_dir(doc, root) / "gates-report.json"


def output_dir(doc: str, root: Path = Path(".")) -> Path:
    return root / "output" / doc


def output_zip(doc: str, root: Path = Path(".")) -> Path:
    return root / "output" / f"{doc}.zip"


def output_file(doc: str, fmt: str, root: Path = Path(".")) -> Path:
    ext = {"html": "html", "reqif": "reqif", "reqifz": "reqifz"}.get(fmt, "md")
    return output_dir(doc, root) / f"{doc}.{ext}"


def okf_index(doc: str, root: Path = Path(".")) -> Path:
    return output_dir(doc, root) / "index.md"


def okf_section(doc: str, index: int, slug: str, root: Path = Path(".")) -> Path:
    return output_dir(doc, root) / f"{index:02d}-{slug}.md"


def assets_dir(doc: str, root: Path = Path(".")) -> Path:
    return output_dir(doc, root) / "assets"


def grade_shards_dir(doc: str, root: Path = Path(".")) -> Path:
    return output_dir(doc, root) / "grade-shards"


def grade_shard_path(doc: str, page_number: int, root: Path = Path(".")) -> Path:
    return grade_shards_dir(doc, root) / f"page{page_number}.json"


def grade_report_json(doc: str, root: Path = Path(".")) -> Path:
    return output_dir(doc, root) / "grade-report.json"
